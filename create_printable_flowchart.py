#!/usr/bin/env python3
"""One-page printable flowchart + update PowerPoint with field readings."""

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.colors import Color, white, black
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Colors
NAVY = Color(0.04, 0.17, 0.29)
TEAL = Color(0.05, 0.49, 0.48)
ORANGE = Color(0.77, 0.36, 0.15)
GREEN = Color(0.18, 0.49, 0.20)
RED = Color(0.72, 0.11, 0.11)
AMBER = Color(0.90, 0.32, 0.00)
LIGHT = Color(0.96, 0.97, 0.98)
LGREEN = Color(0.91, 0.96, 0.91)
LRED = Color(1.0, 0.92, 0.93)
LORANGE = Color(0.99, 0.91, 0.85)
LTEAL = Color(0.83, 0.94, 0.93)
LBLUE = Color(0.89, 0.95, 0.99)
GRAY = Color(0.35, 0.42, 0.48)


def round_rect(c, x, y, w, h, fill, stroke=None, radius=4):
    c.setFillColor(fill)
    if stroke:
        c.setStrokeColor(stroke)
        c.setLineWidth(1.5)
    else:
        c.setStrokeColor(fill)
        c.setLineWidth(0.5)
    c.roundRect(x, y, w, h, radius, fill=1, stroke=1)


def text(c, x, y, s, size=10, color=black, bold=False, align="left", max_w=None):
    c.setFillColor(color)
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    if max_w and align == "center":
        lines = simpleSplit(s, "Helvetica-Bold" if bold else "Helvetica", size, max_w)
        for i, line in enumerate(lines):
            c.drawCentredString(x, y - i * (size + 2), line)
        return len(lines)
    if align == "center":
        c.drawCentredString(x, y, s)
    elif align == "right":
        c.drawRightString(x, y, s)
    else:
        c.drawString(x, y, s)
    return 1


def arrow_down(c, x, y, length=10, color=TEAL):
    c.setFillColor(color)
    c.setStrokeColor(color)
    c.setLineWidth(2)
    c.line(x, y, x, y - length + 4)
    path = c.beginPath()
    path.moveTo(x - 4, y - length + 5)
    path.lineTo(x + 4, y - length + 5)
    path.lineTo(x, y - length)
    path.close()
    c.drawPath(path, fill=1, stroke=0)


def create_pdf(path):
    page = landscape(A4)
    W, H = page
    c = canvas.Canvas(path, pagesize=page)
    m = 12 * mm

    # Background
    c.setFillColor(LIGHT)
    c.rect(0, 0, W, H, fill=1, stroke=0)

    # Header
    c.setFillColor(NAVY)
    c.rect(0, H - 22 * mm, W, 22 * mm, fill=1, stroke=0)
    c.setFillColor(TEAL)
    c.rect(0, H - 24 * mm, W, 2 * mm, fill=1, stroke=0)

    text(c, m, H - 10 * mm, "NO.4 E/R VENT. FAN (REV) — Start / Run / Stop & Fault Flow", 16, white, True)
    text(
        c, m, H - 17 * mm,
        "HI-2536-ME-SD1/SD2  ·  SMC-503H  ·  Printable one-page guide  ·  Field readings included",
        9, Color(0.66, 0.77, 0.85),
    )

    # ===== LEFT: Start-Run-Stop sequence =====
    left_x = m
    left_w = 95 * mm
    y = H - 30 * mm

    round_rect(c, left_x, y - 6 * mm, left_w, 8 * mm, TEAL)
    text(c, left_x + left_w / 2, y - 3.5 * mm, "START → RUN → STOP SEQUENCE", 11, white, True, "center")

    seq = [
        ("1  PRE-START", "52 ON · SOURCE ON · ABNORMAL OFF · Amp ≈ 0", TEAL),
        ("2  PRESS SUPPLY (or EXHAUST)", "SMC 4F/4R ON · HEATING OFF · 88SH open", TEAL),
        ("3  REDUCED-VOLTAGE START", "6-1 / 6-2 + ATr · limit inrush", ORANGE),
        ("4  TIMER 19 ≈ 30 s", "Transfer start → full voltage", ORANGE),
        ("5  FULL-VOLTAGE RUN", "88F or 88R · 4FX/4RX feedback to SMC", GREEN),
        ("6  MONITOR", "Door A · EOCR A · lamps · load", GREEN),
        ("7  NORMAL STOP", "STOP → contactors open → HEATING ON", AMBER),
    ]

    y = y - 10 * mm
    box_h = 11.5 * mm
    for i, (title, body, col) in enumerate(seq):
        round_rect(c, left_x, y - box_h, left_w, box_h, white, col, 3)
        c.setFillColor(col)
        c.roundRect(left_x, y - box_h, 3.5 * mm, box_h, 1, fill=1, stroke=0)
        text(c, left_x + 6 * mm, y - 4.5 * mm, title, 9, NAVY, True)
        text(c, left_x + 6 * mm, y - 9 * mm, body, 8, GRAY)
        if i < len(seq) - 1:
            arrow_down(c, left_x + left_w / 2, y - box_h - 0.5 * mm, 4, col)
            y = y - box_h - 4.5 * mm
        else:
            y = y - box_h

    # ===== RIGHT TOP: Field observations =====
    rx = left_x + left_w + 8 * mm
    rw = W - rx - m
    y2 = H - 30 * mm

    round_rect(c, rx, y2 - 6 * mm, rw, 8 * mm, ORANGE)
    text(c, rx + rw / 2, y2 - 3.5 * mm, "YOUR FIELD OBSERVATIONS (filled)", 11, white, True, "center")

    y2 = y2 - 10 * mm
    obs = [
        ("After stop lamps", "SOURCE ON · ABNORMAL OFF · INTERVAL OFF · HEATING ON", LGREEN, GREEN),
        ("Door ammeter (running)", "~ 63 A   (FLA = 65.4 A) → looks normal", LBLUE, TEAL),
        ("EOCR measured current", "3.7 A secondary  (> SET’G 3.27 A)", LORANGE, ORANGE),
        ("EOCR SET’G (drawing)", "3.27 A  ·  CHA=INV  ·  D-TIME 6 s  ·  O-TIME 10 s", LORANGE, ORANGE),
        ("88F", "Already renewed — stop still occurs", LRED, RED),
        ("Lamp test", "All lamps OK (ABNORMAL lamp works)", LGREEN, GREEN),
        ("Interpretation", "Looks like NORMAL/command stop OR EOCR quiet trip\nResolve 63 A vs EOCR 3.7 A with clamp meter", LORANGE, ORANGE),
    ]

    for title, body, fill, stroke in obs:
        lines = body.count("\n") + 1
        bh = 9 * mm if lines == 1 else 12 * mm
        round_rect(c, rx, y2 - bh, rw, bh, fill, stroke, 3)
        text(c, rx + 3 * mm, y2 - 3.5 * mm, title, 8, stroke, True)
        for i, line in enumerate(body.split("\n")):
            text(c, rx + 3 * mm, y2 - 7 * mm - i * 3.5 * mm, line, 8, black)
        y2 = y2 - bh - 2.5 * mm

    # ===== RIGHT BOTTOM: Decision flow =====
    round_rect(c, rx, y2 - 6 * mm, rw, 8 * mm, NAVY)
    text(c, rx + rw / 2, y2 - 3.5 * mm, "WHEN IT STOPS — DECISION FLOW", 11, white, True, "center")
    y2 = y2 - 10 * mm

    # Two columns of decisions
    col_w = (rw - 4 * mm) / 2
    # Left decision
    round_rect(c, rx, 18 * mm, col_w, y2 - 18 * mm, LRED, RED, 3)
    text(c, rx + col_w / 2, y2 - 5 * mm, "IF ABNORMAL = ON", 9, RED, True, "center")
    abn_lines = [
        "1. Check EOCR for TRIP / OL",
        "2. Clamp all 3 motor phases",
        "3. If clamp high → load / fan mech.",
        "4. If clamp ~63 A but EOCR 3.7",
        "   → CT / EOCR false reading",
        "5. Do NOT only raise SET’G",
        "6. Reset only after cause found",
    ]
    yy = y2 - 10 * mm
    for line in abn_lines:
        text(c, rx + 3 * mm, yy, line, 7.5, black)
        yy -= 3.8 * mm

    # Right decision
    rx2 = rx + col_w + 4 * mm
    round_rect(c, rx2, 18 * mm, col_w, y2 - 18 * mm, LGREEN, GREEN, 3)
    text(c, rx2 + col_w / 2, y2 - 5 * mm, "IF ABNORMAL = OFF (your case)", 9, GREEN, True, "center")
    nor_lines = [
        "1. Note: HEATING ON = clean stop",
        "2. Watch which dies first:",
        "   · SUPPLY lamp first → SMC/STOP/remote",
        "   · Amp→0 while RUN lit → power side",
        "3. Retest in LOCAL only",
        "4. Check STOP NC contact (heat/intermittent)",
        "5. Check SMC 4F seal-in & GSP2 remote",
        "6. Still check EOCR vs clamp (3.7 vs 63)",
    ]
    yy = y2 - 10 * mm
    for line in nor_lines:
        text(c, rx2 + 3 * mm, yy, line, 7.5, black)
        yy -= 3.8 * mm

    # Footer
    text(
        c, m, 8 * mm,
        "CT check: primary ÷ 20 ≈ EOCR secondary (100/5). Example: 63 A ≈ 3.15 A.  |  Max 3 consecutive starts; then wait 40 min for ATr cool-down.",
        7.5, GRAY,
    )
    text(c, W - m, 8 * mm, "Page 1 of 1", 7.5, GRAY, align="right")

    c.showPage()
    c.save()
    print(f"PDF saved: {path}")


def add_pptx_slides(pptx_in, pptx_out):
    """Append printable flowchart summary + field readings slides."""
    prs = Presentation(pptx_in)

    NAVY_R = RGBColor(0x0B, 0x2C, 0x4A)
    TEAL_R = RGBColor(0x0E, 0x7C, 0x7B)
    ORANGE_R = RGBColor(0xC4, 0x5C, 0x26)
    GREEN_R = RGBColor(0x2E, 0x7D, 0x32)
    RED_R = RGBColor(0xB7, 0x1C, 0x1C)
    AMBER_R = RGBColor(0xE6, 0x51, 0x00)
    LIGHT_R = RGBColor(0xF4, 0xF7, 0xFA)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    DARK = RGBColor(0x1A, 0x1A, 0x1A)
    GRAY_R = RGBColor(0x5A, 0x6A, 0x7A)
    LG = RGBColor(0xE8, 0xF5, 0xE9)
    LR = RGBColor(0xFF, 0xEB, 0xEE)
    LO = RGBColor(0xFD, 0xE8, 0xD8)
    LT = RGBColor(0xD4, 0xEF, 0xEE)
    LB = RGBColor(0xE3, 0xF2, 0xFD)

    def bg(slide):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_R
        shape.line.fill.background()
        spTree = slide.shapes._spTree
        sp = shape._element
        spTree.remove(sp)
        spTree.insert(2, sp)

    def header(slide, title, sub):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY_R
        bar.line.fill.background()
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.95), prs.slide_width, Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL_R
        accent.line.fill.background()
        box = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.7))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size = Pt(24)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = "Calibri"
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(0xA8, 0xC5, 0xD8)
        r2.font.name = "Calibri"

    def footer(slide, page):
        box = slide.shapes.add_textbox(Inches(0.4), Inches(7.15), Inches(10), Inches(0.3))
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = "NO.4 E/R VENT. FAN (REV)  |  HI-2536-ME-SD1/SD2  |  Field readings included"
        r.font.size = Pt(10)
        r.font.color.rgb = GRAY_R
        r.font.name = "Calibri"
        box2 = slide.shapes.add_textbox(Inches(11.5), Inches(7.15), Inches(1.5), Inches(0.3))
        p2 = box2.text_frame.paragraphs[0]
        p2.alignment = PP_ALIGN.RIGHT
        r2 = p2.add_run()
        r2.text = f"{page} / 11"
        r2.font.size = Pt(10)
        r2.font.color.rgb = GRAY_R
        r2.font.name = "Calibri"

    def card(slide, left, top, w, h, fill, line):
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        s.adjustments[0] = 0.08
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
        return s

    def tb(slide, left, top, w, h, text, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(left, top, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
        return box

    # Update existing footers page counts roughly by adding new slides as 10, 11
    # --- Slide: Field readings ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    header(slide, "9. Your Field Readings (as found)", "Filled from panel checks during troubleshooting")
    footer(slide, 10)

    rows = [
        ("After-stop lamps", "SOURCE ON · ABNORMAL OFF · INTERVAL OFF · HEATING ON", "Matches NORMAL / command stop state", GREEN_R, LG),
        ("Door ammeter (running)", "~ 63 A  (FLA 65.4 A)", "Under FLA — looks normal", TEAL_R, LB),
        ("EOCR measured", "3.7 A secondary", "Above SET’G 3.27 A — INV may trip in minutes", ORANGE_R, LO),
        ("EOCR setting (drawing)", "SET’G 3.27 A · CHA INV · D-TIME 6 s", "Expected secondary at FLA ≈ 3.27 A (CT 100/5)", ORANGE_R, LO),
        ("88F contactor", "Already renewed", "Did not cure delayed stop", RED_R, LR),
        ("Lamp test", "All indicators OK", "ABNORMAL lamp is good — so OFF means truly no alarm", GREEN_R, LG),
        ("Conflict to resolve", "Door ~63 A vs EOCR 3.7 A", "Use clamp meter on all 3 phases to decide real vs false", AMBER_R, LO),
    ]

    for i, (a, b, ctext, accent, fill) in enumerate(rows):
        top = Inches(1.2 + i * 0.8)
        card(slide, Inches(0.35), top, Inches(12.6), Inches(0.72), fill, accent)
        tb(slide, Inches(0.5), top + Inches(0.08), Inches(3.2), Inches(0.55), a, 12, True, accent)
        tb(slide, Inches(3.7), top + Inches(0.08), Inches(4.2), Inches(0.55), b, 12, True, DARK)
        tb(slide, Inches(8.0), top + Inches(0.08), Inches(4.7), Inches(0.55), ctext, 11, False, DARK)

    # --- Slide: One-page flowchart (PPT version) ---
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    header(slide, "10. One-Page Flowchart (print companion)", "Same content as the A4 landscape PDF — keep beside the panel")
    footer(slide, 11)

    # Three columns
    cols = [
        (
            "START → RUN",
            TEAL_R,
            LT,
            [
                "1. Pre-start OK (52, SOURCE, no abnormal)",
                "2. Press SUPPLY or EXHAUST",
                "3. Heater OFF (88SH open)",
                "4. Reduced V start (6-1/6-2 + ATr)",
                "5. Timer 19 ≈ 30 s transition",
                "6. Full voltage via 88F / 88R",
                "7. 4FX / 4RX feedback to SMC",
                "8. Monitor A + EOCR + lamps",
            ],
        ),
        (
            "NORMAL STOP",
            GREEN_R,
            LG,
            [
                "STOP / command drop",
                "4F/4R release",
                "88F/88R open → Amp 0",
                "HEATING ON (if heater ON)",
                "ABNORMAL stays OFF",
                "",
                "Your after-stop state matched this",
                "→ check STOP / SMC / remote",
            ],
        ),
        (
            "NEXT ACTION",
            ORANGE_R,
            LO,
            [
                "Clamp all 3 phases while running",
                "Compare to door A and EOCR 3.7",
                "",
                "If clamp ≈ 63 A: EOCR/CT suspect",
                "If clamp high: real overload",
                "",
                "Watch stop: SUPPLY lamp first?",
                "Retest in LOCAL only",
            ],
        ),
    ]

    for i, (title, accent, fill, lines) in enumerate(cols):
        left = Inches(0.35 + i * 4.3)
        card(slide, left, Inches(1.25), Inches(4.1), Inches(5.6), fill, accent)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), Inches(1.45), Inches(3.7), Inches(0.5)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = accent
        badge.line.fill.background()
        badge.adjustments[0] = 0.15
        tb(slide, left + Inches(0.2), Inches(1.52), Inches(3.7), Inches(0.4), title, 14, True, WHITE, PP_ALIGN.CENTER)
        body = "\n".join(lines)
        tb(slide, left + Inches(0.25), Inches(2.15), Inches(3.6), Inches(4.4), body, 13, False, DARK)

    prs.save(pptx_out)
    print(f"PPTX saved: {pptx_out}")


if __name__ == "__main__":
    pdf_path = "/workspace/No4_ER_Vent_Fan_Printable_Flowchart.pdf"
    pdf_art = "/opt/cursor/artifacts/No4_ER_Vent_Fan_Printable_Flowchart.pdf"
    pptx_in = "/workspace/No4_ER_Vent_Fan_Start_Run_Stop_Sequence.pptx"
    pptx_out = "/workspace/No4_ER_Vent_Fan_Start_Run_Stop_Sequence.pptx"
    pptx_art = "/opt/cursor/artifacts/No4_ER_Vent_Fan_Start_Run_Stop_Sequence.pptx"

    create_pdf(pdf_path)
    create_pdf(pdf_art)
    add_pptx_slides(pptx_in, pptx_out)
    # copy updated pptx to artifacts
    import shutil
    shutil.copy2(pptx_out, pptx_art)
    print("Done.")
