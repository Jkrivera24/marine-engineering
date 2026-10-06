#!/usr/bin/env python3
"""Generate No.4 E/R Vent Fan start/run/stop sequence PowerPoint."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
NAVY = RGBColor(0x0B, 0x2C, 0x4A)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
ORANGE = RGBColor(0xC4, 0x5C, 0x26)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xB7, 0x1C, 0x1C)
AMBER = RGBColor(0xE6, 0x51, 0x00)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x6A, 0x7A)
LIGHT_TEAL = RGBColor(0xD4, 0xEF, 0xEE)
LIGHT_ORANGE = RGBColor(0xFD, 0xE8, 0xD8)
LIGHT_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
LIGHT_RED = RGBColor(0xFF, 0xEB, 0xEE)
LIGHT_BLUE = RGBColor(0xE3, 0xF2, 0xFD)


def add_bg(slide, color=LIGHT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_header_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.95)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(0.95), prs.slide_width, Inches(0.08)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Calibri"
    if subtitle:
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(12)
        run2.font.color.rgb = RGBColor(0xA8, 0xC5, 0xD8)
        run2.font.name = "Calibri"


def add_footer(slide, page, total=10):
    box = slide.shapes.add_textbox(
        Inches(0.4), Inches(7.15), Inches(10), Inches(0.3)
    )
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "NO.4 E/R VENT. FAN (REV)  |  HI-2536-ME-SD1/SD2  |  Hyundai Heavy Industries"
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY
    run.font.name = "Calibri"

    box2 = slide.shapes.add_textbox(
        Inches(11.5), Inches(7.15), Inches(1.5), Inches(0.3)
    )
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    run2 = p2.add_run()
    run2.text = f"{page} / {total}"
    run2.font.size = Pt(10)
    run2.font.color.rgb = GRAY
    run2.font.name = "Calibri"


def add_textbox(slide, left, top, width, height, text, size=14, bold=False,
                color=DARK, align=PP_ALIGN.LEFT, font_name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return box


def add_card(slide, left, top, width, height, fill=WHITE, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.adjustments[0] = 0.08
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.color.rgb = RGBColor(0xD0, 0xD8, 0xE0)
        shape.line.width = Pt(1)
    return shape


def add_step_box(slide, left, top, width, height, number, title, body, accent=TEAL):
    card = add_card(slide, left, top, width, height, WHITE, accent)

    num = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left + Inches(0.15), top + Inches(0.15), Inches(0.4), Inches(0.4)
    )
    num.fill.solid()
    num.fill.fore_color.rgb = accent
    num.line.fill.background()

    nbox = slide.shapes.add_textbox(
        left + Inches(0.15), top + Inches(0.2), Inches(0.4), Inches(0.35)
    )
    tf = nbox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(number)
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Calibri"

    add_textbox(
        slide, left + Inches(0.65), top + Inches(0.15), width - Inches(0.8), Inches(0.35),
        title, size=13, bold=True, color=NAVY
    )
    add_textbox(
        slide, left + Inches(0.15), top + Inches(0.6), width - Inches(0.3), height - Inches(0.7),
        body, size=11, color=DARK
    )
    return card


def add_arrow_right(slide, left, top, color=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, left, top, Inches(0.35), Inches(0.25)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_arrow_down(slide, left, top, color=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, left, top, Inches(0.28), Inches(0.32)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


TOTAL = 9

# ========== SLIDE 1: Title ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)

band = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, 0, Inches(2.2), prs.slide_width, Inches(3.0)
)
band.fill.solid()
band.fill.fore_color.rgb = RGBColor(0x0E, 0x3A, 0x5C)
band.line.fill.background()

add_textbox(
    slide, Inches(0.8), Inches(2.4), Inches(11.5), Inches(0.5),
    "NO.4 E/R VENT. FAN (REV)", size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER
)
add_textbox(
    slide, Inches(0.8), Inches(3.05), Inches(11.5), Inches(0.45),
    "Start → Run → Stop Sequence Guide", size=24, color=RGBColor(0x7E, 0xD6, 0xD4),
    align=PP_ALIGN.CENTER
)
add_textbox(
    slide, Inches(0.8), Inches(3.65), Inches(11.5), Inches(0.8),
    "Based on schematic HI-2536-ME-SD1 (p.072) & HI-2536-ME-SD2 (p.073)\n"
    "Controller: LUXCO SMC-503H  |  Autotransformer reduced-voltage start  |  EOCR 51",
    size=14, color=RGBColor(0xA8, 0xC5, 0xD8), align=PP_ALIGN.CENTER
)

add_textbox(
    slide, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4),
    "Hyundai Heavy Industries  ·  Shipboard Engine Room Ventilation Fan (Reversible)",
    size=12, color=RGBColor(0x7A, 0x9A, 0xB0), align=PP_ALIGN.CENTER
)

# ========== SLIDE 2: System overview ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "1. System Overview", "Main components involved in start / run / stop")
add_footer(slide, 2, TOTAL)

items = [
    ("52", "Main MCCB", "3-phase power in\n(R, S, T)", TEAL),
    ("88F / 88R", "Direction Contactors", "Forward (Supply)\nReverse (Exhaust)", TEAL),
    ("ATr + 6-1/6-2", "Reduced-Voltage Start", "Autotransformer\n80% / 60% taps", ORANGE),
    ("19", "Transition Timer", "≈ 30 sec\nStart → Full voltage", ORANGE),
    ("51 EOCR", "Overcurrent Relay", "SET 3.27 A\nINV characteristic", RED),
    ("SMC-503H", "Motor Controller", "Local/Remote F-R\nLamps & alarms", GREEN),
    ("4FX / 4RX", "Run Feedback", "Aux contacts back\nto SMC", GREEN),
    ("88SH", "Space Heater", "ON when motor\nstopped (if enabled)", AMBER),
]

for i, (code, title, body, color) in enumerate(items):
    col = i % 4
    row = i // 4
    left = Inches(0.4 + col * 3.2)
    top = Inches(1.3 + row * 2.7)
    add_card(slide, left, top, Inches(3.0), Inches(2.4), WHITE, color)
    # code badge
    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(0.25),
        Inches(2.6), Inches(0.45)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()
    badge.adjustments[0] = 0.2
    add_textbox(
        slide, left + Inches(0.2), top + Inches(0.3), Inches(2.6), Inches(0.4),
        code, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.2), top + Inches(0.9), Inches(2.6), Inches(0.4),
        title, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.2), top + Inches(1.35), Inches(2.6), Inches(0.85),
        body, size=12, color=DARK, align=PP_ALIGN.CENTER
    )

# ========== SLIDE 3: Pre-start conditions ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "2. Pre-Start Conditions (Ready to Start)", "All must be OK before pressing SUPPLY or EXHAUST")
add_footer(slide, 3, TOTAL)

checks = [
    ("1", "MCCB 52 CLOSED", "3-phase power available to starter panel"),
    ("2", "SOURCE lamp ON", "SMC-503H has AC 220 V control power (S1–S2)"),
    ("3", "ABNORMAL lamp OFF", "No fault latched; EOCR not tripped"),
    ("4", "HEATING may be ON", "Normal if HEATER switch = ON and motor was stopped"),
    ("5", "Ammeter ≈ 0 A", "Motor not running; contactors open"),
    ("6", "EOCR healthy", "Not in TRIP; SET’G = 3.27 A (drawing value)"),
    ("7", "Direction clear", "Wait until fan stopped mechanically before F↔R change"),
    ("8", "Start count OK", "Max 3 consecutive starts; then wait 40 min (ATr cool-down)"),
]

for i, (n, title, body) in enumerate(checks):
    col = i % 2
    row = i // 2
    left = Inches(0.4 + col * 6.4)
    top = Inches(1.25 + row * 1.35)
    add_card(slide, left, top, Inches(6.15), Inches(1.2), WHITE, TEAL)
    num = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left + Inches(0.2), top + Inches(0.35), Inches(0.5), Inches(0.5)
    )
    num.fill.solid()
    num.fill.fore_color.rgb = TEAL
    num.line.fill.background()
    add_textbox(
        slide, left + Inches(0.2), top + Inches(0.42), Inches(0.5), Inches(0.4),
        n, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.9), top + Inches(0.2), Inches(5.0), Inches(0.4),
        title, size=15, bold=True, color=NAVY
    )
    add_textbox(
        slide, left + Inches(0.9), top + Inches(0.6), Inches(5.0), Inches(0.45),
        body, size=12, color=DARK
    )

# ========== SLIDE 4: Start sequence overview ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "3. Start Sequence — Overview (Supply / Forward)", "Same logic for Exhaust / Reverse using 88R & 4RX")
add_footer(slide, 4, TOTAL)

steps = [
    ("Press\nSUPPLY", "Operator\ncommand", TEAL),
    ("4F / SMC\noutput", "Run command\nto starter", TEAL),
    ("88SH\nOFF", "Heater drops\nout", AMBER),
    ("Start\n6-1 / 6-2", "Reduced V\nvia ATr", ORANGE),
    ("88F\ncloses", "Forward\npower path", GREEN),
    ("Timer 19\n≈ 30 s", "Transition\nwait", ORANGE),
    ("Full\nvoltage", "Run contactors\ntransfer", GREEN),
    ("4FX\nfeedback", "SMC confirms\nRUN", GREEN),
]

for i, (title, body, color) in enumerate(steps):
    left = Inches(0.25 + i * 1.62)
    top = Inches(1.5)
    add_card(slide, left, top, Inches(1.5), Inches(2.6), WHITE, color)
    circ = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left + Inches(0.45), top + Inches(0.2), Inches(0.55), Inches(0.55)
    )
    circ.fill.solid()
    circ.fill.fore_color.rgb = color
    circ.line.fill.background()
    add_textbox(
        slide, left + Inches(0.45), top + Inches(0.3), Inches(0.55), Inches(0.4),
        str(i + 1), size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.08), top + Inches(0.9), Inches(1.35), Inches(0.9),
        title, size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.08), top + Inches(1.85), Inches(1.35), Inches(0.6),
        body, size=10, color=GRAY, align=PP_ALIGN.CENTER
    )
    if i < len(steps) - 1:
        add_arrow_right(slide, left + Inches(1.48), top + Inches(1.15), color)

# Expected lamps
add_card(slide, Inches(0.4), Inches(4.4), Inches(12.5), Inches(2.4), LIGHT_GREEN, GREEN)
add_textbox(
    slide, Inches(0.6), Inches(4.55), Inches(12), Inches(0.35),
    "Expected indications after successful start (running)",
    size=14, bold=True, color=GREEN
)
add_textbox(
    slide, Inches(0.6), Inches(5.05), Inches(12), Inches(1.5),
    "• SOURCE = ON (white)\n"
    "• SUPPLY (or EXHAUST) RUN = ON (green)\n"
    "• ABNORMAL = OFF   |   INTERVAL = OFF   |   HEATING = OFF\n"
    "• Ammeter ≈ motor load (typical ~60–65 A near FLA 65.4 A)\n"
    "• EOCR measured current ≈ 3.0–3.3 A secondary (CT 100/5) under normal load",
    size=13, color=DARK
)

# ========== SLIDE 5: Detailed start sequence ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "4. Start Sequence — Detailed Steps", "Power & control actions in order")
add_footer(slide, 5, TOTAL)

detail = [
    ("A", "Command", TEAL,
     "Operator presses SUPPLY (F) on SMC-503H (or remote F via GSP2).\n"
     "Internal 4F picks up. HEATING lamp goes OFF; 88SH drops."),
    ("B", "Reduced-voltage start", ORANGE,
     "Start contactors 6-1 / 6-2 close. Motor fed through autotransformer (ATr)\n"
     "at reduced voltage (80%/60% taps) to limit inrush. 88F (or 88R) path active."),
    ("C", "Transition (Timer 19)", ORANGE,
     "Timer 19 runs ≈ 30 seconds. After time-out, circuit transfers from\n"
     "reduced-voltage start to full-voltage run. 6-1/6-2 sequence completes."),
    ("D", "Full-voltage run", GREEN,
     "Motor runs on full line voltage via 88F (Supply) or 88R (Exhaust).\n"
     "Interlocks prevent 88F and 88R both closed. Hour meter (HM) counts."),
    ("E", "Feedback & protect", GREEN,
     "4FX (or 4RX) aux contact closes → SMC sees RUN feedback.\n"
     "EOCR 51 monitors CT current continuously (INV curve, D-TIME 6 s start delay)."),
]

for i, (letter, title, color, body) in enumerate(detail):
    top = Inches(1.2 + i * 1.1)
    add_card(slide, Inches(0.4), top, Inches(12.5), Inches(1.0), WHITE, color)
    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), top + Inches(0.25),
        Inches(1.6), Inches(0.5)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()
    badge.adjustments[0] = 0.15
    add_textbox(
        slide, Inches(0.55), top + Inches(0.32), Inches(1.6), Inches(0.4),
        f"{letter}. {title}", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, Inches(2.4), top + Inches(0.2), Inches(10.2), Inches(0.7),
        body, size=12, color=DARK
    )

# ========== SLIDE 6: Running monitoring ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "5. Running / Monitoring", "What to watch while the fan is running")
add_footer(slide, 6, TOTAL)

mon = [
    ("Door ammeter", "~60–65 A typical\nFLA = 65.4 A", "If rising toward/\nabove FLA → load issue", LIGHT_BLUE, TEAL),
    ("EOCR display", "Measured A (sec.)\nSET’G = 3.27 A", "If measured > SET’G\n→ INV trip in minutes", LIGHT_ORANGE, ORANGE),
    ("SMC lamps", "SOURCE + RUN ON\nABNORMAL OFF", "ABNORMAL ON =\nfault / overload path", LIGHT_GREEN, GREEN),
    ("Feedback", "4FX / 4RX closed\nto SMC", "Lost feedback may\ncause abnormal stop", LIGHT_TEAL, TEAL),
]

for i, (title, mid, tip, fill, accent) in enumerate(mon):
    left = Inches(0.35 + i * 3.25)
    add_card(slide, left, Inches(1.3), Inches(3.1), Inches(3.5), fill, accent)
    add_textbox(
        slide, left + Inches(0.15), Inches(1.5), Inches(2.8), Inches(0.45),
        title, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.15), Inches(2.15), Inches(2.8), Inches(1.1),
        mid, size=13, color=DARK, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.15), Inches(3.5), Inches(2.8), Inches(1.0),
        tip, size=12, bold=True, color=accent, align=PP_ALIGN.CENTER
    )

# EOCR table
add_card(slide, Inches(0.4), Inches(5.05), Inches(12.5), Inches(1.85), WHITE, RED)
add_textbox(
    slide, Inches(0.6), Inches(5.15), Inches(12), Inches(0.35),
    "EOCR 51 settings (from drawing)", size=14, bold=True, color=RED
)
add_textbox(
    slide, Inches(0.6), Inches(5.55), Inches(12), Inches(1.2),
    "Motor rating 65.4 A   |   Range 0.5–6 A   |   SET’G 3.27 A   |   O-TIME 10 s   |   D-TIME 6 s   |   CHA = INV   |   PF = OFF\n"
    "CT ratio (effective): 100/5 → secondary = primary ÷ 20   (example: 63 A primary ≈ 3.15 A on EOCR)\n"
    "Important: If EOCR measured > 3.27 while door meter looks normal, verify with clamp meter — EOCR may trip on INV delay.",
    size=12, color=DARK
)

# ========== SLIDE 7: Normal stop ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "6. Normal Stop Sequence", "Clean stop — no fault declared")
add_footer(slide, 7, TOTAL)

stop_steps = [
    ("1", "Stop command", "Press STOP on SMC\nor remote opens run\ncommand", TEAL),
    ("2", "4F / 4R drop", "SMC releases\nforward/reverse\noutput", TEAL),
    ("3", "88F / 88R open", "Power contactors\nopen; motor coasts\ndown", ORANGE),
    ("4", "Current → 0", "Ammeter returns\nto 0 A", ORANGE),
    ("5", "Heater ON", "If HEATER = ON,\n88SH closes;\nHEATING lamp ON", AMBER),
    ("6", "Final state", "SOURCE ON\nABNORMAL OFF\nRUN lamps OFF", GREEN),
]

for i, (n, title, body, color) in enumerate(stop_steps):
    left = Inches(0.3 + i * 2.15)
    add_card(slide, left, Inches(1.35), Inches(2.05), Inches(3.3), WHITE, color)
    circ = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left + Inches(0.7), Inches(1.55), Inches(0.55), Inches(0.55)
    )
    circ.fill.solid()
    circ.fill.fore_color.rgb = color
    circ.line.fill.background()
    add_textbox(
        slide, left + Inches(0.7), Inches(1.65), Inches(0.55), Inches(0.4),
        n, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.1), Inches(2.3), Inches(1.85), Inches(0.5),
        title, size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER
    )
    add_textbox(
        slide, left + Inches(0.1), Inches(2.9), Inches(1.85), Inches(1.5),
        body, size=11, color=DARK, align=PP_ALIGN.CENTER
    )
    if i < len(stop_steps) - 1:
        add_arrow_right(slide, left + Inches(1.95), Inches(2.7), color)

add_card(slide, Inches(0.4), Inches(4.95), Inches(12.5), Inches(1.9), LIGHT_ORANGE, ORANGE)
add_textbox(
    slide, Inches(0.6), Inches(5.1), Inches(12), Inches(0.35),
    "Your observed after-stop state matched NORMAL STOP",
    size=14, bold=True, color=ORANGE
)
add_textbox(
    slide, Inches(0.6), Inches(5.55), Inches(12), Inches(1.1),
    "SOURCE ON · ABNORMAL OFF · INTERVAL OFF · HEATING ON · Ammeter 0 A\n"
    "That means SMC released the run command (or control path opened as a “quiet” stop),\n"
    "not a declared abnormal trip. If unintended: check STOP contact, SMC seal-in (4F), remote GSP2.",
    size=13, color=DARK
)

# ========== SLIDE 8: Abnormal vs normal ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "7. Normal Stop vs Abnormal Stop", "How to tell what stopped the fan")
add_footer(slide, 8, TOTAL)

add_card(slide, Inches(0.35), Inches(1.25), Inches(6.2), Inches(5.5), LIGHT_GREEN, GREEN)
add_textbox(
    slide, Inches(0.55), Inches(1.4), Inches(5.8), Inches(0.45),
    "NORMAL / COMMAND STOP", size=18, bold=True, color=GREEN, align=PP_ALIGN.CENTER
)
add_textbox(
    slide, Inches(0.55), Inches(2.0), Inches(5.8), Inches(4.4),
    "Lamps\n"
    "• ABNORMAL = OFF\n"
    "• HEATING = ON (if heater switch ON)\n"
    "• RUN lamps = OFF\n"
    "• INTERVAL = OFF\n\n"
    "Typical causes\n"
    "• STOP button pressed / intermittent NC\n"
    "• SMC 4F seal-in dropped\n"
    "• Remote run contact opened (GSP2)\n"
    "• Control power blip (may not latch abnormal)\n\n"
    "88F replacement usually does NOT fix this",
    size=13, color=DARK
)

add_card(slide, Inches(6.8), Inches(1.25), Inches(6.2), Inches(5.5), LIGHT_RED, RED)
add_textbox(
    slide, Inches(7.0), Inches(1.4), Inches(5.8), Inches(0.45),
    "ABNORMAL / PROTECTION STOP", size=18, bold=True, color=RED, align=PP_ALIGN.CENTER
)
add_textbox(
    slide, Inches(7.0), Inches(2.0), Inches(5.8), Inches(4.4),
    "Lamps\n"
    "• ABNORMAL = ON (red)\n"
    "• EOCR may show TRIP / OL\n"
    "• RUN lamps = OFF\n"
    "• HEATING may still come ON after stop\n\n"
    "Typical causes\n"
    "• EOCR 51 trip (INV — can take minutes)\n"
    "• Overload / high current\n"
    "• Power fail path\n"
    "• Lost start/run feedback (4FX) after T1\n\n"
    "Check EOCR trip state + clamp current\nbefore raising SET’G",
    size=13, color=DARK
)

# ========== SLIDE 9: Quick reference flowchart ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_header_bar(slide, "8. Quick Troubleshooting Flow (Stops After Several Minutes)", "Use after-stop lamps + current readings together")
add_footer(slide, 9, TOTAL)

# Flow boxes
flow = [
    (Inches(4.5), Inches(1.2), "Fan stops after several minutes", NAVY, WHITE),
]

# decision diamond-ish using rounded rects
boxes = [
    (0.4, 2.0, 4.0, 1.3, "ABNORMAL ON?", "Yes → EOCR / protect\nCheck EOCR TRIP, clamp A", RED, LIGHT_RED),
    (6.9, 2.0, 5.9, 1.3, "ABNORMAL OFF?", "Yes → Command stop path\nSTOP / SMC 4F / Remote", GREEN, LIGHT_GREEN),
    (0.4, 3.6, 4.0, 1.4, "EOCR measured > 3.27?", "Even if door amp ~63 A\nVerify with clamp meter", ORANGE, LIGHT_ORANGE),
    (6.9, 3.6, 5.9, 1.4, "SUPPLY lamp off first?", "Yes = SMC dropped command\nNo = power side after SMC", TEAL, LIGHT_TEAL),
    (0.4, 5.3, 6.0, 1.4, "Clamp ≈ EOCR (high)", "Real overload / mechanical load\nDo not only raise SET’G", RED, LIGHT_RED),
    (6.8, 5.3, 6.0, 1.4, "Clamp ≈ door (~63), EOCR high", "False EOCR/CT reading\nCheck CT & EOCR device", ORANGE, LIGHT_ORANGE),
]

add_card(slide, Inches(3.5), Inches(1.2), Inches(6.3), Inches(0.65), NAVY, NAVY)
add_textbox(
    slide, Inches(3.5), Inches(1.3), Inches(6.3), Inches(0.5),
    "Fan stops after several minutes — check lamps at the moment of stop",
    size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER
)

for left, top, w, h, title, body, accent, fill in boxes:
    add_card(slide, Inches(left), Inches(top), Inches(w), Inches(h), fill, accent)
    add_textbox(
        slide, Inches(left + 0.15), Inches(top + 0.15), Inches(w - 0.3), Inches(0.35),
        title, size=13, bold=True, color=accent
    )
    add_textbox(
        slide, Inches(left + 0.15), Inches(top + 0.55), Inches(w - 0.3), Inches(0.7),
        body, size=12, color=DARK
    )

# Save
out1 = "/workspace/No4_ER_Vent_Fan_Start_Run_Stop_Sequence.pptx"
out2 = "/opt/cursor/artifacts/No4_ER_Vent_Fan_Start_Run_Stop_Sequence.pptx"
prs.save(out1)
prs.save(out2)
print(f"Saved: {out1}")
print(f"Saved: {out2}")
